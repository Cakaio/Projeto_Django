"""O kit de papel do Bazar: o plano B do evento.

Existe porque é a ÚNICA camada que funciona sem servidor, sem rede e sem
bateria. Impresso na sexta, vale no sábado aconteça o que acontecer — e num
evento que acontece uma vez por ano, sem ensaio, essa é a diferença entre
atrasar a fila e perder o dia.

Duas folhas, e a divisão não é estética:

  FICHA DO CAIXA — uma linha por CRIANÇA, não por peça. À mão, com fila, linha
  por peça é lentidão garantida. Traz a tabela de pontos no cabeçalho porque
  quem soma de cabeça precisa dela na frente.

  LISTA DE ELEGÍVEIS POR SALINHA — substitui a busca quando não há busca. Traz
  idade e numerações ao lado do nome, o que resolve homônimo no papel do mesmo
  jeito que resolve na tela.

A montagem vive aqui, longe de HTTP, para poder ser testada sem requisição e
para servir tanto o HTML de impressão quanto o .xlsx.
"""
from atendido.models import LISTA_SALAS, Atendido

from .models import Categoria, SalaDoBazar

# Quantas linhas em branco a ficha do caixa traz. Doze cabe numa folha A4 com
# a tabela de pontos em cima e ainda sobra margem para escrever torto.
LINHAS_EM_BRANCO = 12


def idade_em_anos(atendido, hoje):
    """Idade da criança, ou None se a ficha não tiver nascimento."""
    nascimento = getattr(atendido, "data_nascimento", None)
    if not nascimento:
        return None
    return hoje.year - nascimento.year - (
        (hoje.month, hoje.day) < (nascimento.month, nascimento.day))


def folhas_do_kit(bazar, hoje=None):
    """Tudo que as duas folhas precisam, numa estrutura só.

    Salinha sem ninguém NÃO vira folha: papel gasto numa véspera corrida é
    papel que some antes do sábado.
    """
    from django.utils import timezone

    hoje = hoje or timezone.localdate()
    rotulos = dict(LISTA_SALAS)

    categorias = list(
        Categoria.objects.filter(bazar=bazar, ativo=True).order_by("ordem", "nome"))
    salas = list(SalaDoBazar.objects.filter(bazar=bazar, ativo=True))

    por_salinha = []
    for codigo, rotulo in LISTA_SALAS:
        pessoas = list(
            Atendido.objects.ativos().filter(sala=codigo).order_by("nome"))
        if not pessoas:
            continue
        por_salinha.append({
            "sala": codigo,
            "rotulo": rotulos.get(codigo, codigo),
            "atendidos": [
                {
                    "nome": pessoa.nome,
                    "idade": idade_em_anos(pessoa, hoje),
                    "camisa": pessoa.numeracao_camisa or "",
                    "calca": pessoa.numeracao_calca or "",
                    "calcado": pessoa.numeracao_calcado or "",
                }
                for pessoa in pessoas
            ],
        })

    return {
        "bazar": bazar,
        "categorias": categorias,
        "salas": salas,
        "por_salinha": por_salinha,
        "linhas_em_branco": range(LINHAS_EM_BRANCO),
        "total_elegiveis": sum(len(g["atendidos"]) for g in por_salinha),
    }


def planilha_do_kit(bazar, hoje=None):
    """O mesmo conteúdo em .xlsx, para quem prefere digitar no Excel depois.

    Import local de propósito: `openpyxl` só é carregado por quem pediu a
    planilha, e não na subida do projeto.
    """
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    folhas = folhas_do_kit(bazar, hoje)
    livro = Workbook()

    cabecalho = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="E8560F")
    centro = Alignment(horizontal="center")

    # ── Aba 1: ficha do caixa ──
    ficha = livro.active
    ficha.title = "Ficha do caixa"
    ficha["A1"] = f"{bazar.nome} — {bazar.data:%d/%m/%Y}"
    ficha["A1"].font = Font(bold=True, size=14)
    ficha["A2"] = f"Cota da 1ª etapa: {bazar.cota_inicial} pontos por criança"
    ficha["A3"] = "Pontos por categoria: " + " · ".join(
        f"{categoria.nome} = {categoria.pontos}" for categoria in folhas["categorias"])

    colunas = (["Nº", "Nome do atendido", "Salinha"]
               + [categoria.nome for categoria in folhas["categorias"]]
               + ["Total de pontos", "Quem levou", "Hora", "Conferiu"])
    ficha.append([])
    ficha.append(colunas)
    for celula in ficha[5]:
        celula.font = cabecalho
        celula.fill = fundo
        celula.alignment = centro
    for numero in range(1, LINHAS_EM_BRANCO + 1):
        ficha.append([numero] + [""] * (len(colunas) - 1))

    ficha.column_dimensions["B"].width = 30
    ficha.column_dimensions["C"].width = 16

    # ── Aba 2: elegíveis ──
    elegiveis = livro.create_sheet("Elegíveis")
    elegiveis.append(["Salinha", "Nome", "Idade", "Camisa", "Calça", "Calçado",
                      "Passou?"])
    for celula in elegiveis[1]:
        celula.font = cabecalho
        celula.fill = fundo
    for grupo in folhas["por_salinha"]:
        for pessoa in grupo["atendidos"]:
            elegiveis.append([
                grupo["rotulo"], pessoa["nome"], pessoa["idade"],
                pessoa["camisa"], pessoa["calca"], pessoa["calcado"], "",
            ])
    elegiveis.column_dimensions["A"].width = 18
    elegiveis.column_dimensions["B"].width = 32

    fluxo = io.BytesIO()
    livro.save(fluxo)
    return fluxo.getvalue()
