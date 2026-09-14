"""O kit de papel: o plano B do evento.

É a única camada que funciona sem servidor, sem rede e sem bateria — e por
isso é a que menos pode depender de qualquer coisa. Aberto a QUALQUER
voluntário logado: hoje o único CSV está atrás da coordenação, o que é inútil,
porque quem está no caixa é que precisa imprimir.
"""
from datetime import date

from django.test import RequestFactory, TestCase

from atendido.models import Atendido
from voluntario.models import Voluntario

from . import views
from .models import Bazar, Categoria, SalaDoBazar


class KitDePapelTest(TestCase):
    def setUp(self):
        self.fabrica = RequestFactory()
        self.bazar = Bazar.objects.create(
            nome="Bazar 2026", data=date(2026, 10, 3),
            etapa=Bazar.Etapa.PRIMEIRA, cota_inicial=5)
        Categoria.objects.create(bazar=self.bazar, nome="Camiseta",
                                 pontos=1, ordem=1)
        Categoria.objects.create(bazar=self.bazar, nome="Calçado",
                                 pontos=3, ordem=2)
        Categoria.objects.create(bazar=self.bazar, nome="Saiu de linha",
                                 pontos=2, ordem=3, ativo=False)
        SalaDoBazar.objects.create(bazar=self.bazar, nome="Sala 1", ordem=1)

        self.joao = Atendido.objects.create(
            nome="João Silva", sala="VIOLETA", data_nascimento=date(2015, 1, 1),
            numeracao_camisa="10")
        self.voluntario = Voluntario.objects.create_user(
            username="vol_kit", password="pw", area="AMARELO")

    def _html(self):
        requisicao = self.fabrica.get(f"/bazar/{self.bazar.pk}/kit/")
        requisicao.user = self.voluntario
        return views.kit_papel(requisicao, pk=self.bazar.pk).content.decode()

    # ── quem alcança ──

    def test_qualquer_voluntario_logado_imprime(self):
        requisicao = self.fabrica.get(f"/bazar/{self.bazar.pk}/kit/")
        requisicao.user = self.voluntario
        self.assertEqual(
            views.kit_papel(requisicao, pk=self.bazar.pk).status_code, 200)

    # ── a ficha do caixa ──

    def test_traz_a_tabela_de_pontos(self):
        """Quem soma à mão precisa da tabela de preços na frente."""
        html = self._html()
        self.assertIn("Camiseta", html)
        self.assertIn("Calçado", html)

    def test_categoria_desativada_fica_de_fora(self):
        """Imprimir categoria que não existe mais faz o voluntário oferecer."""
        self.assertNotIn("Saiu de linha", self._html())

    def test_traz_a_cota_em_destaque(self):
        self.assertIn("5", self._html())

    def test_tem_linhas_em_branco_para_preencher(self):
        """Uma linha por CRIANÇA, não por peça: à mão, linha por peça é
        lentidão garantida."""
        html = self._html()
        self.assertGreaterEqual(html.count('class="linha-branco"'), 10)

    # ── a lista de elegíveis ──

    def test_lista_os_elegiveis_com_idade_e_numeracao(self):
        """Substitui a busca quando não há busca — e resolve homônimo no papel
        do mesmo jeito que na tela."""
        html = self._html()
        self.assertIn("João Silva", html)
        self.assertIn("10", html)

    # ── folha de impressão de verdade ──

    def test_nao_estende_base_html(self):
        """Folha de impressão não carrega navbar, sidebar nem JS do site."""
        html = self._html()
        self.assertNotIn("pcf-navitem", html)
        self.assertNotIn("bazar-atendimento.js", html)

    def test_o_botao_de_imprimir_some_na_impressao(self):
        html = self._html()
        self.assertIn("@media print", html)
        self.assertIn("print-color-adjust", html)

    # ── planilha ──

    def test_a_planilha_sai_em_xlsx(self):
        requisicao = self.fabrica.get(f"/bazar/{self.bazar.pk}/kit.xlsx")
        requisicao.user = self.voluntario
        resposta = views.kit_planilha(requisicao, pk=self.bazar.pk)

        self.assertEqual(resposta.status_code, 200)
        self.assertIn("spreadsheetml", resposta["Content-Type"])
        # .xlsx é um zip: começa com "PK".
        self.assertTrue(resposta.content.startswith(b"PK"))

    def test_a_planilha_tem_as_duas_abas(self):
        import io
        from openpyxl import load_workbook

        requisicao = self.fabrica.get(f"/bazar/{self.bazar.pk}/kit.xlsx")
        requisicao.user = self.voluntario
        conteudo = views.kit_planilha(requisicao, pk=self.bazar.pk).content

        planilha = load_workbook(io.BytesIO(conteudo))
        self.assertIn("Ficha do caixa", planilha.sheetnames)
        self.assertIn("Elegíveis", planilha.sheetnames)


class FolhasDoKitTest(TestCase):
    """A montagem, longe de HTTP."""

    def setUp(self):
        self.bazar = Bazar.objects.create(
            nome="Bazar 2026", data=date(2026, 10, 3), cota_inicial=5)
        Categoria.objects.create(bazar=self.bazar, nome="Camiseta", pontos=1)
        Atendido.objects.create(nome="Ana", sala="AMARELO",
                                data_nascimento=date(2016, 5, 2))
        Atendido.objects.create(nome="Bruno", sala="VIOLETA",
                                data_nascimento=date(2014, 2, 2))

    def test_agrupa_os_elegiveis_por_salinha(self):
        from .papelaria import folhas_do_kit
        folhas = folhas_do_kit(self.bazar)
        salas = {grupo["rotulo"]: grupo["atendidos"] for grupo in folhas["por_salinha"]}
        self.assertEqual(len(salas["Amarelo"]), 1)
        self.assertEqual(len(salas["Violeta"]), 1)

    def test_salinha_sem_ninguem_nao_vira_folha(self):
        """Imprimir folha vazia é papel gasto numa véspera corrida."""
        from .papelaria import folhas_do_kit
        rotulos = [g["rotulo"] for g in folhas_do_kit(self.bazar)["por_salinha"]]
        self.assertEqual(sorted(rotulos), ["Amarelo", "Violeta"])
